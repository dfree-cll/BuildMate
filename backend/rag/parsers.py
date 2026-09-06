"""Deterministic document extraction for RAG ingestion.

Geometry never becomes a RAG fact source. IFC extraction is limited to textual
names/types/properties useful for document lookup; geometric queries still use
Model IR and geometry engines.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from backend.domain.errors import DependencyFailure, ValidationFailure


@dataclass(frozen=True)
class ParsedKnowledgeDocument:
    content: str
    parser_version: str
    metadata: dict


def parse_knowledge_file(path: Path, *, max_chars: int = 5_000_000) -> ParsedKnowledgeDocument:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        content = path.read_text(encoding="utf-8", errors="replace")
        parser = "text-1"
        metadata = {}
    elif suffix == ".pdf":
        content, metadata = _parse_pdf(path)
        parser = "pymupdf-1"
    elif suffix == ".docx":
        content, metadata = _parse_docx(path)
        parser = "python-docx-1"
    elif suffix == ".ifc":
        content, metadata = _parse_ifc_text(path)
        parser = "ifcopenshell-text-1"
    else:
        raise ValidationFailure(f"artifact is not a supported knowledge document: {suffix}")
    content = content.strip()
    if not content:
        raise ValidationFailure("document parser produced no text; OCR/manual text is required")
    if len(content) > max_chars:
        raise ValidationFailure(f"parsed document exceeds {max_chars} characters")
    return ParsedKnowledgeDocument(content=content, parser_version=parser, metadata=metadata)


def _parse_pdf(path: Path) -> tuple[str, dict]:
    try:
        import fitz
    except ImportError as exc:
        raise DependencyFailure("PyMuPDF is required for PDF ingestion") from exc
    try:
        document = fitz.open(path)
        pages = []
        for index, page in enumerate(document, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append(f"# Page {index}\n\n{text}")
        page_count = document.page_count
        document.close()
    except (OSError, RuntimeError, ValueError) as exc:
        raise ValidationFailure(f"invalid PDF: {str(exc)[:300]}") from exc
    ocr_used = False
    if not pages:
        # Reuse the existing bounded RapidOCR pipeline for scanned PDFs.  The
        # fallback remains explicit: if OCR is unavailable or produces no
        # text, ingestion stops instead of indexing an empty/fake document.
        try:
            from backend.engines.pdf_parser import _sync_extract_pdf

            fallback = _sync_extract_pdf(str(path))
        except Exception as exc:
            raise DependencyFailure("OCR pipeline is unavailable for scanned PDF") from exc
        raw_text = str(fallback.get("raw_text") or "").strip()
        if raw_text:
            sections = [part.strip() for part in raw_text.split("\n\n---PAGE BREAK---\n\n") if part.strip()]
            pages = [f"# Page {index}\n\n{section}" for index, section in enumerate(sections, start=1)]
            ocr_used = bool(pages)
    if not pages:
        raise ValidationFailure("PDF has no extractable text; OCR produced no text")
    return "\n\n".join(pages), {"page_count": page_count, "ocr_used": ocr_used}


def _parse_docx(path: Path) -> tuple[str, dict]:
    try:
        from docx import Document
    except ImportError as exc:
        raise DependencyFailure("python-docx is required for DOCX ingestion") from exc
    try:
        document = Document(path)
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    paragraphs.append(" | ".join(values))
    except Exception as exc:
        raise ValidationFailure(f"invalid DOCX: {str(exc)[:300]}") from exc
    return "\n\n".join(paragraphs), {"paragraph_count": len(paragraphs)}


def _parse_ifc_text(path: Path) -> tuple[str, dict]:
    try:
        import ifcopenshell
    except ImportError as exc:
        raise DependencyFailure("IfcOpenShell is required for IFC text ingestion") from exc
    try:
        model = ifcopenshell.open(str(path))
        products = model.by_type("IfcProduct")[:20_000]
    except Exception as exc:
        raise ValidationFailure(f"invalid IFC: {str(exc)[:300]}") from exc
    lines = []
    for product in products:
        global_id = getattr(product, "GlobalId", "") or ""
        name = getattr(product, "Name", "") or ""
        description = getattr(product, "Description", "") or ""
        lines.append(f"{product.is_a()} | {global_id} | {name} | {description}")
    return "\n".join(lines), {"ifc_product_count": len(products), "geometry_authority": False}
